#!/usr/bin/env python3
"""
SuperCat Customer File Converter
=================================
Converts your customer Excel files into the format required for SuperCat/eCat import.

HOW TO USE:
  1. Make sure Python 3 is installed (https://www.python.org/downloads/)
  2. Install the required library by running this in your terminal:
         pip install openpyxl
  3. Run this script:
         python convert_customers_to_supercat.py
  4. A file picker will open — select one or more .xlsx customer files
  5. Converted files will be saved in the same folder with " - SUPERCAT" added to the name

EXPECTED INPUT COLUMNS (your file should have these):
  Customer Number, Customer Name, Address 1, Address 2,
  City, State, Zip, COUNTRY, Phone 1, Salesperson ID

WHAT THE SCRIPT DOES:
  - Renames columns to match SuperCat's import format
  - Adds a DefaultPriceCode column (set to "dn" for Dealer Price)
  - Adds blank Ship-To address columns for you to fill in
  - Saves the result as a new Excel file (your original is NOT modified)

If you have questions, reach out to Kylor at SuperCat support.
"""

import sys
import os
from pathlib import Path

COLUMN_MAP = {
    "Customer Number": "BillToCode",
    "Customer Name":   "BillToName",
    "Address 1":       "BillToAddress1",
    "Address 2":       "BillToAddress2",
    "City":            "BillToCity",
    "State":           "BillToState",
    "Zip":             "BillToPostCode",
    "COUNTRY":         "BillToCountry",
    "Phone 1":         "BuyerPhone",
    "Salesperson ID":  "TerritoryCodes",
}

ADDED_COLUMNS = {
    "DefaultPriceCode": "dn",
    "ShipToAddress1":   "",
    "ShipToCity":       "",
    "ShipToState":      "",
    "ShipToPostCode":   "",
}


def check_dependencies():
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        print("\n[!] Missing required library: openpyxl")
        print("    Install it by running:  pip install openpyxl\n")
        sys.exit(1)


def pick_files():
    """Open a file dialog and return selected file paths."""
    try:
        import tkinter as tk
        from tkinter import filedialog
        root = tk.Tk()
        root.withdraw()
        root.attributes("-topmost", True)
        paths = filedialog.askopenfilenames(
            title="Select your customer Excel file(s)",
            filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")],
        )
        root.destroy()
        return list(paths)
    except Exception:
        return None


def convert_file(filepath):
    """Convert a single customer file to SuperCat format. Returns output path."""
    from openpyxl import load_workbook, Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = load_workbook(filepath)
    ws = wb.active

    headers = [cell.value for cell in ws[1]]

    missing = [col for col in COLUMN_MAP if col not in headers]
    if missing:
        raise ValueError(
            f"Missing expected columns: {', '.join(missing)}\n"
            f"  Found columns: {', '.join(str(h) for h in headers)}\n"
            f"  Make sure your file has: {', '.join(COLUMN_MAP.keys())}"
        )

    col_indices = {name: headers.index(name) for name in COLUMN_MAP}

    out_wb = Workbook()
    out_ws = out_wb.active
    out_ws.title = "Customers"

    out_headers = list(COLUMN_MAP.values()) + list(ADDED_COLUMNS.keys())

    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_text = Font(bold=True, color="FFFFFF")

    for col_idx, header in enumerate(out_headers, start=1):
        cell = out_ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_text
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    data_row_count = 0
    for row_num in range(2, ws.max_row + 1):
        row_data = [cell.value for cell in ws[row_num]]
        if all(v is None for v in row_data):
            continue

        out_row = []
        for orig_col in COLUMN_MAP:
            val = row_data[col_indices[orig_col]]
            if val is not None:
                val = str(val).strip()
            else:
                val = ""
            out_row.append(val)

        for default_val in ADDED_COLUMNS.values():
            out_row.append(default_val)

        for col_idx, val in enumerate(out_row, start=1):
            out_ws.cell(row=row_num, column=col_idx, value=val if val != "" else None)

        data_row_count += 1

    for col in out_ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        out_ws.column_dimensions[col_letter].width = min(max_len + 4, 40)

    p = Path(filepath)
    out_name = f"{p.stem} - SUPERCAT{p.suffix}"
    out_path = p.parent / out_name
    out_wb.save(out_path)

    return str(out_path), data_row_count


def main():
    print("=" * 56)
    print("  SuperCat Customer File Converter")
    print("=" * 56)
    print()

    check_dependencies()

    files = None

    if len(sys.argv) > 1:
        files = [f for f in sys.argv[1:] if f.endswith(".xlsx")]
        if not files:
            print("[!] No .xlsx files found in arguments.")
            print("    Usage: python convert_customers_to_supercat.py file1.xlsx file2.xlsx")
            sys.exit(1)
    else:
        print("Opening file picker...")
        files = pick_files()

    if not files:
        print("\nNo files selected. You can also pass files as arguments:")
        print("  python convert_customers_to_supercat.py \"MY CUSTOMER LIST.xlsx\"")
        sys.exit(0)

    print(f"\nProcessing {len(files)} file(s)...\n")

    success_count = 0
    for filepath in files:
        filename = os.path.basename(filepath)
        print(f"  [{filename}]")
        try:
            out_path, row_count = convert_file(filepath)
            print(f"    ✓ Converted {row_count} customers")
            print(f"    → Saved to: {os.path.basename(out_path)}")
            success_count += 1
        except ValueError as e:
            print(f"    ✗ Error: {e}")
        except Exception as e:
            print(f"    ✗ Unexpected error: {e}")
        print()

    print("-" * 56)
    print(f"  Done! {success_count}/{len(files)} file(s) converted.")
    print()
    print("  NEXT STEPS:")
    print("  1. Open the new file(s) ending in ' - SUPERCAT'")
    print("  2. Fill in the Ship-To address columns")
    print("     (can be same as billing if orders ship there)")
    print("  3. Send the completed file(s) back to SuperCat")
    print("-" * 56)

    if sys.platform == "win32":
        input("\nPress Enter to close...")


if __name__ == "__main__":
    main()
