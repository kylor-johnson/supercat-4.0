import openpyxl
import csv
import re

ML_PATH = '/Users/kylorjohnson/Downloads/Attachments (2)/ML CUSTOMER LIST REVISED - UPDATED.xlsx'
NSL_PATH = '/Users/kylorjohnson/Downloads/Attachments (2)/NSL CUSTOMER LIST REVISED - UPDATED.xlsx'
OUTPUT_PATH = '/Users/kylorjohnson/Downloads/customers.csv'

OUTPUT_HEADERS = [
    'BillToCode', 'BillToName', 'BillToAddress1', 'BillToAddress2',
    'BillToCity', 'BillToState', 'BillToPostCode', 'BillToCountry',
    'BuyerPhone', 'TerritoryCodes', 'DefaultPriceCode',
    'ShipToAddress1', 'ShipToCity', 'ShipToState', 'ShipToPostCode',
    'ShipToEmail',
]

EMAIL_RE = re.compile(r'^[^@\s]+@[^@\s]+\.[^@\s]+$')


def read_xlsx(path):
    wb = openpyxl.load_workbook(path)
    ws = wb['Customers']
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        d = {h: (str(v).strip() if v is not None else '') for h, v in zip(headers, row)}
        rows.append(d)
    return rows


def process_rows(rows):
    emails_moved = 0
    for row in rows:
        row.setdefault('ShipToEmail', '')
        addr = row.get('BillToAddress1', '').strip()
        if addr and EMAIL_RE.match(addr):
            row['ShipToEmail'] = addr
            row['BillToAddress1'] = ''
            emails_moved += 1
    return emails_moved


ml_rows = read_xlsx(ML_PATH)
nsl_rows = read_xlsx(NSL_PATH)

print(f'ML rows read: {len(ml_rows)}')
print(f'NSL rows read: {len(nsl_rows)}')

ml_emails = process_rows(ml_rows)
nsl_emails = process_rows(nsl_rows)
print(f'Emails moved from BillToAddress1 -> ShipToEmail: ML={ml_emails}, NSL={nsl_emails}')

for row in ml_rows:
    row['DefaultPriceCode'] = 'mldn'
for row in nsl_rows:
    row['DefaultPriceCode'] = 'nsldn'

combined = ml_rows + nsl_rows
combined.sort(key=lambda r: r.get('BillToCode', ''))

print(f'Combined rows: {len(combined)}')

with open(OUTPUT_PATH, 'w', newline='', encoding='utf-8') as f:
    writer = csv.DictWriter(f, fieldnames=OUTPUT_HEADERS, extrasaction='ignore')
    writer.writeheader()
    writer.writerows(combined)

print(f'Written to {OUTPUT_PATH}')
